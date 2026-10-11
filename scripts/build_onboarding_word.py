from __future__ import annotations

import argparse
import re
import tempfile
from collections.abc import Iterable
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Hiragino Sans GB"
LATIN_FONT = "Arial"
CODE_FONT = "Menlo"
NAVY = "1F4E78"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "D9D9D9"
CODE_BG = "F3F4F6"
BLACK = "000000"
WHITE = "FFFFFF"


def set_run_font(run, name: str, size: float | None = None, *, bold: bool | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN_FONT if name == BODY_FONT else name)
    fonts.set(qn("w:hAnsi"), LATIN_FONT if name == BODY_FONT else name)
    fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "5")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), LIGHT_GRAY)


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_hyperlink(paragraph, text: str, target: str, *, internal: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if internal:
        hyperlink.set(qn("w:anchor"), target)
    else:
        relationship_id = paragraph.part.relate_to(
            target,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    properties.extend([fonts, color, underline])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, BODY_FONT, 8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), BODY_FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    field_properties.extend([fonts, size])
    field_run.append(field_properties)
    value = OxmlElement("w:t")
    value.text = "1"
    field_run.append(value)
    field.append(field_run)
    paragraph._p.append(field)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, BODY_FONT, 8.5)


def sanitize_heading(text: str) -> str:
    text = re.sub(r"[`*_]", "", text)
    text = re.sub(r"^(\d+)\.(\d+)\s*", r"第 \1 章 第 \2 节 ", text)
    text = re.sub(r"^(\d+)\.\s*", r"第 \1 章 ", text)
    text = re.sub(r"[：:？?！!。；;，,、｜|（）()“”‘’\[\]{}<>]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


INLINE_PATTERN = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
)


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, BODY_FONT, 11)
        token = match.group(0)
        if token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, BODY_FONT, 11, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, CODE_FONT, 9.5)
            run.font.color.rgb = RGBColor(54, 65, 83)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), CODE_BG)
            run._element.get_or_add_rPr().append(shading)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, BODY_FONT, 11)
            run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, BODY_FONT, 11)


def add_body_paragraph(document: Document, text: str, *, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    add_inline(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.28
    return paragraph


def add_quote(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.2
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(64, 64, 64)
    return paragraph


def add_code_block(document: Document, lines: list[str], language: str, diagram_path: Path) -> None:
    if language == "mermaid":
        create_architecture_diagram(diagram_path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        shape = run.add_picture(str(diagram_path), width=Inches(6.55))
        properties = shape._inline.docPr
        properties.set("descr", "ExperimentLens 系统架构图")
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(8)
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_BG)
    properties.append(shading)
    for index, line in enumerate(lines or [""]):
        run = paragraph.add_run(line)
        set_run_font(run, CODE_FONT, 8.7)
        if index < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(cells: Iterable[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    if len(normalized) > 1 and is_separator_row(normalized[1]):
        normalized.pop(1)
    table = document.add_table(rows=len(normalized), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    total_width = 6.85
    weights = []
    for column in range(column_count):
        longest = max(len(row[column]) for row in normalized)
        weights.append(max(7, min(34, longest)))
    weight_sum = sum(weights)
    widths = [total_width * weight / weight_sum for weight in weights]
    for row_index, row in enumerate(normalized):
        prevent_row_split(table.rows[row_index])
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, BODY_FONT, 9.2, bold=row_index == 0)
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    repeat_table_header(table.rows[0])
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def create_architecture_diagram(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1800, 1040), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    font = ImageFont.truetype(font_path, 34)
    small = ImageFont.truetype(font_path, 27)

    def box(x: int, y: int, width: int, height: int, text: str, fill: str = "#EEF4F8"):
        draw.rounded_rectangle(
            (x, y, x + width, y + height), radius=18, fill=fill, outline="#1F4E78", width=4
        )
        bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=8, align="center")
        tx = x + (width - (bbox[2] - bbox[0])) / 2
        ty = y + (height - (bbox[3] - bbox[1])) / 2
        draw.multiline_text((tx, ty), text, font=font, fill="#111111", spacing=8, align="center")

    def arrow(x1: int, y1: int, x2: int, y2: int):
        draw.line((x1, y1, x2, y2), fill="#5B6B7A", width=6)
        angle = __import__("math").atan2(y2 - y1, x2 - x1)
        length = 20
        spread = 0.55
        points = [
            (x2, y2),
            (x2 - length * __import__("math").cos(angle - spread), y2 - length * __import__("math").sin(angle - spread)),
            (x2 - length * __import__("math").cos(angle + spread), y2 - length * __import__("math").sin(angle + spread)),
        ]
        draw.polygon(points, fill="#5B6B7A")

    box(80, 70, 330, 115, "用户与 React 工作台")
    box(550, 70, 330, 115, "FastAPI 控制面")
    box(1020, 70, 330, 115, "SQLite 或 PostgreSQL")
    arrow(410, 127, 550, 127)
    arrow(880, 127, 1020, 127)

    box(550, 300, 330, 115, "独立 Agent Worker", "#E8F1FA")
    arrow(1185, 185, 1185, 260)
    arrow(1185, 260, 880, 340)

    box(80, 520, 330, 125, "Model Gateway\n离线模型或 OpenAI")
    box(550, 520, 330, 125, "六个只读工具")
    box(1020, 520, 330, 125, "项目工作区\nParquet 与 JSON")
    arrow(650, 415, 360, 520)
    arrow(715, 415, 715, 520)
    arrow(880, 582, 1020, 582)

    box(80, 780, 330, 125, "Evidence Validator")
    box(550, 780, 330, 125, "SSE 事件与 Trace")
    box(1020, 780, 330, 125, "Evaluation Executor")
    arrow(650, 645, 360, 780)
    arrow(715, 645, 715, 780)
    arrow(1185, 780, 830, 645)

    draw.text((1450, 80), "主链路", font=small, fill="#1F4E78")
    draw.text((1450, 130), "控制  执行  证据  评测", font=small, fill="#555555")
    image.save(path, format="PNG", optimize=True)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = BODY_FONT
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.paragraph_format.space_after = Pt(18)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 17, 12, 8),
        ("Heading 2", 14, 10, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def parse_source(source: str) -> tuple[list[tuple[str, str]], list[str]]:
    sections: list[tuple[str, str]] = []
    lines = source.splitlines()
    for line in lines:
        if line.startswith("## "):
            sections.append((sanitize_heading(line[3:]), f"section_{len(sections) + 1}"))
    return sections, lines


def build_document(source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8")
    sections, lines = parse_source(source)
    document = Document()
    configure_styles(document)
    document.core_properties.title = "ExperimentLens 从零学习与面试讲解手册"
    document.core_properties.subject = "AI Agent 初学者学习与项目接管指南"
    document.core_properties.author = "ExperimentLens Project"
    document.core_properties.keywords = "AI Agent, ExperimentLens, 面试, 学习手册"

    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    add_page_number(section.footer.paragraphs[0])

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(76)
    title = document.add_paragraph("ExperimentLens 从零学习与面试讲解手册", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("AI Agent 初学者学习与项目接管指南")
    set_run_font(subtitle_run, BODY_FONT, 14, bold=True)
    subtitle.paragraph_format.space_after = Pt(34)
    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience_run = audience.add_run("适用对象  AI Agent 初学者  项目接手者  求职面试准备者")
    set_run_font(audience_run, BODY_FONT, 11)
    audience.paragraph_format.space_after = Pt(14)
    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline_run = tagline.add_run("Evidence driven AI agent for reliable experiment diagnostics")
    set_run_font(tagline_run, LATIN_FONT, 10)
    tagline_run.italic = True
    document.add_page_break()

    intro = document.add_paragraph("阅读说明", style="Heading 1")
    intro.paragraph_format.page_break_before = False
    add_body_paragraph(
        document,
        "本手册帮助第一次接触 AI Agent 的读者理解 ExperimentLens 的业务价值、数据链路、Agent Runtime 和工程取舍，并逐步具备运行、修改、验证和讲解项目的能力。",
    )
    add_body_paragraph(
        document,
        "阅读时请同时对照项目代码、测试和评测报告。面试表述必须基于真实参与和可验证贡献，不应把尚未实现的能力或他人的工作描述为自己独立完成。",
    )
    add_body_paragraph(
        document,
        "建议先完成第 1 至 5 章建立直觉，再学习架构与技术栈，最后按照第 16 章进行七天实操并用第 19 章自测。",
    )
    toc = document.add_paragraph("目录", style="Heading 1")
    toc.paragraph_format.page_break_before = True
    for heading, bookmark in sections:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        add_hyperlink(paragraph, heading, bookmark, internal=True)
    document.add_page_break()

    diagram_path = Path(tempfile.gettempdir()) / "experimentlens-architecture-word.png"
    first_section = True
    section_number = 0
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    started_body = False
    quote_lines: list[str] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = []

    def flush_quote() -> None:
        nonlocal quote_lines
        if quote_lines:
            add_quote(document, " ".join(line for line in quote_lines if line))
            quote_lines = []

    for raw_line in lines:
        line = raw_line.rstrip()
        if line.startswith("# ") and not started_body:
            continue
        if line.startswith("## "):
            flush_table()
            flush_quote()
            started_body = True
            section_number += 1
            heading = sanitize_heading(line[3:])
            paragraph = document.add_paragraph(heading, style="Heading 1")
            paragraph.paragraph_format.page_break_before = not first_section
            add_bookmark(paragraph, f"section_{section_number}", section_number)
            first_section = False
            continue
        if not started_body:
            continue
        if line.startswith("```"):
            flush_table()
            flush_quote()
            if not in_code:
                in_code = True
                code_language = line[3:].strip()
                code_lines = []
            else:
                add_code_block(document, code_lines, code_language, diagram_path)
                in_code = False
                code_language = ""
                code_lines = []
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            flush_quote()
            table_rows.append(parse_table_row(line))
            continue
        flush_table()
        if line.startswith(">"):
            quote_lines.append(line[1:].strip())
            continue
        flush_quote()
        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("### "):
            document.add_paragraph(sanitize_heading(line[4:]), style="Heading 2")
        elif line.startswith("#### "):
            document.add_paragraph(sanitize_heading(line[5:]), style="Heading 3")
        elif re.match(r"^- \[[ xX]\] ", line):
            checked = line[3].lower() == "x"
            add_body_paragraph(document, f"{'☑' if checked else '☐'} {line[6:]}")
        elif line.startswith("- "):
            add_body_paragraph(document, line[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            paragraph = add_body_paragraph(document, line)
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        else:
            add_body_paragraph(document, line)
    flush_table()
    flush_quote()
    if in_code:
        add_code_block(document, code_lines, code_language, diagram_path)

    for current_section in document.sections:
        current_section.page_width = Inches(8.5)
        current_section.page_height = Inches(11)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the ExperimentLens onboarding Word guide")
    parser.add_argument(
        "--source", default="docs/onboarding-and-interview-study-guide.md", type=Path
    )
    parser.add_argument(
        "--output", default="docs/ExperimentLens-Agent-学习与面试手册.docx", type=Path
    )
    args = parser.parse_args()
    build_document(args.source, args.output)


if __name__ == "__main__":
    main()
